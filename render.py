# -*- coding: utf-8 -*-
"""
Render deck .pptx từ file nội dung YAML/JSON.

    python render.py content.yaml out/deck.pptx

Không cần file template.pptx: mọi layout được dựng bằng code từ theme.py,
nên đổi màu/font chỉ cần sửa theme.py là toàn bộ deck đổi theo.
"""
import sys, io, json, os

# Console Windows mặc định cp1252 -> ép UTF-8 để in được tiếng Việt
for _s in ("stdout", "stderr"):
    _st = getattr(sys, _s, None)
    if _st and hasattr(_st, "buffer"):
        setattr(sys, _s, io.TextIOWrapper(_st.buffer, encoding="utf-8",
                                          errors="replace", line_buffering=True))

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

import theme as T


# ========================= tiện ích dựng hình =========================

def _c(hexstr):
    return RGBColor.from_string(hexstr)


def _blank(prs):
    """Slide trắng hoàn toàn - ta tự vẽ mọi thứ lên."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = _c(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = _c(line)
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def _text(slide, x, y, w, h, txt, size=None, color=T.GRAY, bold=False,
          align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True, spacing=1.0):
    """Đặt một khối chữ. txt có thể là str hoặc list[str] (mỗi phần tử 1 đoạn)."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0

    lines = txt if isinstance(txt, (list, tuple)) else [txt]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(line)
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(4)
        for r in p.runs:
            r.font.size = size or T.SZ_BODY
            r.font.color.rgb = _c(color)
            r.font.bold = bold
            r.font.name = T.FONT
    return box


def _bullets(slide, x, y, w, h, items, size=None, color=T.GRAY, marker="•"):
    """Danh sách gạch đầu dòng. Phần tử có thể là str, hoặc dict{text, level}."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

    for i, it in enumerate(items):
        if isinstance(it, dict):
            txt, lvl = it.get("text", ""), int(it.get("level", 0))
        else:
            txt, lvl = str(it), 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pre = "" if lvl > 0 and marker == "" else (marker + " " if marker else "")
        p.text = ("    " * lvl) + pre + txt
        p.line_spacing = 1.25
        p.space_after = Pt(7)
        for r in p.runs:
            r.font.size = size or T.SZ_BODY
            r.font.color.rgb = _c(color)
            r.font.name = T.FONT
    return box


def _title(slide, text, sub=None):
    """Tiêu đề slide + thanh vàng gạch chân - dùng cho mọi layout nội dung."""
    _text(slide, T.M_LEFT, T.M_TOP, T.SLIDE_W - T.M_LEFT * 2, Pt(44),
          text, size=T.SZ_TITLE, color=T.RED, bold=True)
    _rect(slide, T.M_LEFT, Pt(84), Pt(56), Pt(4), fill=T.YELLOW)
    if sub:
        _text(slide, T.M_LEFT, Pt(93), T.SLIDE_W - T.M_LEFT * 2, Pt(26),
              sub, size=T.SZ_SMALL, color=T.GRAY)


_FONT_CACHE = {}


def _wrap_lines(text, width, fs, bold=False):
    """
    Số dòng chữ sẽ chiếm sau khi xuống dòng trong khung rộng `width`.
    Đo bề rộng bằng chính font Calibri nếu có; không có thì ước lượng.
    """
    key = (fs, bold)
    if key not in _FONT_CACHE:
        f = None
        try:
            from PIL import ImageFont
            for n in (("calibrib.ttf", "arialbd.ttf") if bold
                      else ("calibri.ttf", "arial.ttf")):
                try:
                    f = ImageFont.truetype("C:/Windows/Fonts/" + n, int(fs * 4))
                    break
                except OSError:
                    continue
        except ImportError:
            pass
        _FONT_CACHE[key] = f
    font = _FONT_CACHE[key]

    maxw = width / Pt(1)
    if font is None:                       # không có font -> ước lượng thô
        cpl = max(int(maxw / (fs * 0.52)), 6)
        return max(1, -(-len(text) // cpl))

    # font dựng ở cỡ fs*4 -> chia lại cho đúng tỉ lệ
    def w(t):
        return font.getlength(t) / 4.0

    lines, cur = 1, ""
    for word in text.split():
        t = (cur + " " + word).strip()
        if w(t) <= maxw or not cur:
            cur = t
        else:
            lines += 1
            cur = word
    return lines


def _hero_bg(slide, path=None):
    """
    Ảnh nền trang trí phủ kín slide bìa / slide chương.
    Ảnh đã canh sẵn theo khổ 16:9 nên phủ nguyên khung, không cắt xén.
    Trả về True nếu có vẽ - để layout biết mà thu chữ về nửa trái.
    """
    p = path if path is not None else T.HERO_BG
    if not p or not os.path.exists(p):
        return False
    slide.shapes.add_picture(p, 0, 0, T.SLIDE_W, T.SLIDE_H)
    return True


def _logo_mark(slide):
    """Logo nhỏ đóng góc trên phải của slide nội dung."""
    p = T.LOGO_MARK
    if not p or not os.path.exists(p):
        return
    from PIL import Image as _Im
    iw, ih = _Im.open(p).size
    h = T.LOGO_MARK_H
    w = int(iw * (h / ih))
    slide.shapes.add_picture(p, T.SLIDE_W - T.M_LEFT - w, T.M_TOP, w, h)


def _footer(slide, idx, total):
    """Số trang góc dưới phải."""
    _text(slide, T.SLIDE_W - Pt(90), T.SLIDE_H - Pt(34), Pt(50), Pt(20),
          str(idx), size=T.sz(11), color=T.BORDER, align=PP_ALIGN.RIGHT)


def _numbubble(slide, x, y, n, d=Pt(26), fill=None):
    """Bong bóng số tròn - đặc trưng của deck hướng dẫn phần mềm."""
    s = _rect(slide, x, y, d, d, fill=fill or T.RED, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = str(n)
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        # chữ co theo đường kính bong bóng để luôn nằm gọn bên trong
        r.font.size = min(T.SZ_NUM, Pt(int(d / Pt(1) * 0.58)))
        r.font.bold = True
        r.font.color.rgb = _c(T.WHITE)
        r.font.name = T.FONT
    return s


def _image(slide, path, x, y, maxw, maxh, placeholder_label="Ảnh chụp màn hình"):
    """
    Chèn ảnh vừa khung, giữ tỉ lệ, canh giữa.
    Thiếu ảnh -> vẽ ô gạch chờ để bạn biết chỗ nào cần chụp bổ sung.
    """
    if path and os.path.exists(path):
        from PIL import Image as _Im
        try:
            iw, ih = _Im.open(path).size
        except Exception:
            iw, ih = (16, 9)
        scale = min(maxw / iw, maxh / ih)
        w, h = int(iw * scale), int(ih * scale)
        pic = slide.shapes.add_picture(path, x + (maxw - w) // 2,
                                       y + (maxh - h) // 2, w, h)
        pic.line.color.rgb = _c(T.BORDER)
        pic.line.width = Pt(1)
        return pic

    # --- chưa có ảnh: ô chờ ---
    box = _rect(slide, x, y, maxw, maxh, fill=T.LIGHT, line=T.BORDER)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = f"[ {placeholder_label} ]\n{path or 'chưa gán đường dẫn'}"
    p.alignment = PP_ALIGN.CENTER
    for r in p.runs:
        r.font.size = T.SZ_SMALL
        r.font.color.rgb = _c(T.BORDER)
        r.font.name = T.FONT
    return box


# ============================ CÁC LAYOUT ============================

def lay_title(prs, s):
    """L1 - Slide bìa / kết thúc."""
    sl = _blank(prs)
    _rect(sl, 0, 0, T.SLIDE_W, T.SLIDE_H, fill=T.RED_DARK)
    # "logo": false trong yaml -> slide này không dùng ảnh nền
    has_bg = _hero_bg(sl, s.get("logo")) if s.get("logo") is not False else False
    _rect(sl, 0, T.SLIDE_H - Pt(14), T.SLIDE_W, Pt(14), fill=T.YELLOW)

    # Hoạ tiết nằm nửa phải -> thu chữ về nửa trái cho khỏi bị che
    # Logo loe rộng phía trên, hẹp dần xuống dưới: mép trái của nét đi từ
    # ~500pt (y=180) tới ~447pt (đáy). Chữ dừng ở 430pt là an toàn cả cột.
    tw = Pt(430) - T.M_LEFT if has_bg else T.SLIDE_W - T.M_LEFT * 2
    # Đo bề rộng chữ thật để biết đúng số dòng, rồi canh giữa cả cụm theo
    # chiều dọc - tiêu đề dài hay ngắn đều cân, phụ đề luôn nằm sát bên dưới.
    fs = T.SZ_TITLE_HERO / Pt(1)
    ufs = T.SZ_HERO_SUB / Pt(1)

    th = Pt(fs * 1.2) * _wrap_lines(s["title"], tw, fs, bold=True)
    sub = s.get("subtitle")
    uh = Pt(ufs * 1.35) * _wrap_lines(sub, tw, ufs) if sub else Pt(0)
    gap = Pt(26) if sub else Pt(0)

    top = (T.SLIDE_H - (th + gap + uh)) / 2 - Pt(10)
    _text(sl, T.M_LEFT, top, tw, th + Pt(10),
          s["title"], size=T.SZ_TITLE_HERO, color=T.WHITE, bold=True, spacing=1.15)
    if sub:
        _text(sl, T.M_LEFT, top + th + gap, tw, uh + Pt(10),
              sub, size=T.SZ_HERO_SUB, color=T.YELLOW, spacing=1.25)
    return sl


def lay_section(prs, s):
    """L4 - Slide phân chương: số chương lớn + tên chương."""
    sl = _blank(prs)
    _rect(sl, 0, 0, T.SLIDE_W, T.SLIDE_H, fill=T.RED_DARK)
    # mặc định slide chương KHÔNG dùng ảnh nền; bật bằng "logo": true
    has_bg = _hero_bg(sl) if s.get("logo") is True else False
    # Logo loe rộng phía trên, hẹp dần xuống dưới: mép trái của nét đi từ
    # ~500pt (y=180) tới ~447pt (đáy). Chữ dừng ở 430pt là an toàn cả cột.
    tw = Pt(430) - T.M_LEFT if has_bg else T.SLIDE_W - T.M_LEFT * 2

    # Tính chiều cao từng khối rồi canh giữa cả cụm theo chiều dọc,
    # để slide không bị lệch lên trên khi tên chương chỉ có một dòng.
    nfs = T.SZ_SECTION_NUM / Pt(1)
    sfs = T.SZ_SECTION / Pt(1)
    ufs = T.SZ_SECTION_SUB / Pt(1)

    num = str(s["number"]) if s.get("number") else None
    nh = Pt(nfs * 1.0) if num else Pt(0)
    gap_n = Pt(16) if num else Pt(0)

    sn = _wrap_lines(s["title"], tw, sfs, bold=True)
    th = Pt(sfs * 1.2) * sn

    sub = s.get("subtitle")
    un = _wrap_lines(sub, tw, ufs) if sub else 0
    uh = Pt(ufs * 1.35) * un if sub else Pt(0)
    gap_u = Pt(22) if sub else Pt(0)

    total = nh + gap_n + th + gap_u + uh
    y = (T.SLIDE_H - total) / 2 - Pt(10)      # nhích lên chút cho thoáng đáy

    if num:
        _text(sl, T.M_LEFT, y, Pt(300), nh + Pt(20),
              num, size=T.SZ_SECTION_NUM, color=T.YELLOW, bold=True, spacing=1.0)
        y += nh + gap_n

    _text(sl, T.M_LEFT, y, tw, th + Pt(10),
          s["title"], size=T.SZ_SECTION, color=T.WHITE, bold=True, spacing=1.15)
    y += th + gap_u

    if sub:
        _text(sl, T.M_LEFT, y, tw, uh + Pt(10),
              sub, size=T.SZ_SECTION_SUB, color=T.YELLOW, spacing=1.25)
    return sl


def lay_agenda(prs, s):
    """L3 - Mục lục / danh sách đánh số. Tự chia 2 cột khi > 6 mục."""
    sl = _blank(prs)
    _title(sl, s["title"], s.get("subtitle"))
    items = s.get("items", [])
    two = len(items) > 6
    colw = (T.SLIDE_W - T.M_LEFT * 2 - Pt(40)) / 2 if two else T.SLIDE_W - T.M_LEFT * 2
    per = (len(items) + 1) // 2 if two else len(items)

    for i, it in enumerate(items):
        col, row = (i // per, i % per) if two else (0, i)
        x = T.M_LEFT + col * (colw + Pt(40))
        y = T.CONTENT_TOP + Pt(10) + row * Pt(56)
        _numbubble(sl, x, y, i + 1, d=Pt(30))
        _text(sl, x + Pt(46), y + Pt(3), colw - Pt(46), Pt(48),
              it, size=T.SZ_BODY, color=T.GRAY)
    return sl


def lay_two_col(prs, s):
    """L2 - Hai cột song song (vd: lợi ích phía A / phía B)."""
    sl = _blank(prs)
    _title(sl, s["title"], s.get("subtitle"))
    cols = s.get("columns", [])[:2]
    gap = Pt(32)
    w = (T.SLIDE_W - T.M_LEFT * 2 - gap) / 2
    for i, col in enumerate(cols):
        x = T.M_LEFT + i * (w + gap)
        _rect(sl, x, T.CONTENT_TOP, w, Pt(340), fill=T.LIGHT)
        _rect(sl, x, T.CONTENT_TOP, w, Pt(38), fill=T.RED)
        _text(sl, x + Pt(16), T.CONTENT_TOP + Pt(9), w - Pt(32), Pt(26),
              col.get("heading", ""), size=T.sz(15), color=T.WHITE, bold=True)
        _bullets(sl, x + Pt(16), T.CONTENT_TOP + Pt(54), w - Pt(32), Pt(270),
                 col.get("items", []), size=T.SZ_CALLOUT)
    return sl


def lay_screenshot_callout(prs, s):
    """
    L5 - LAYOUT CHỦ ĐẠO: ảnh chụp màn hình bên trái + callout đánh số bên phải.
    Chiếm >60% deck hướng dẫn phần mềm.
    """
    sl = _blank(prs)
    _title(sl, s["title"], s.get("subtitle"))
    cos = s.get("callouts", [])
    # nhiều callout -> thu hẹp ảnh để chữ có chỗ thở
    imgw = Pt(530) if len(cos) <= 5 else Pt(470)
    imgh = Pt(348)
    _image(sl, s.get("image"), T.M_LEFT, T.CONTENT_TOP, imgw, imgh)

    cx = T.M_LEFT + imgw + Pt(28)
    cw = T.SLIDE_W - cx - T.M_LEFT
    avail = T.SLIDE_H - T.CONTENT_TOP - Pt(34)   # chừa chỗ số trang

    # Ước lượng chiều cao cần, rồi co lại nếu tràn đáy. Ký tự vừa một dòng
    # tỉ lệ nghịch với cỡ chữ, nên tính theo cỡ chữ thực tế đang dùng.
    def measure(fs, gap_line, gap_block):
        tw_ = cw - Pt(max(fs + 11, 22)) - Pt(12)      # trừ bong bóng số
        tot = Pt(0)
        for co in cos:
            t = co.get("text", "") if isinstance(co, dict) else str(co)
            nt = co.get("note") if isinstance(co, dict) else None
            tot += gap_block + gap_line * (_wrap_lines(t, tw_, fs) - 1)
            if nt:
                tot += Pt(28) + Pt(17) * (_wrap_lines(nt, tw_, fs - 2) - 1)
        return tot

    fs = T.SZ_CALLOUT / Pt(1)      # cỡ chữ callout, đơn vị pt
    gl, gb = Pt(19), Pt(34)
    while measure(fs, gl, gb) > avail and fs > 10:
        fs -= 1                     # nhỏ dần 1pt cho đến khi vừa slide
        gl, gb = Pt(fs + 4), Pt(fs + 18)

    size = Pt(fs)
    bub = Pt(max(fs + 11, 22))
    y = T.CONTENT_TOP + Pt(4)
    for i, co in enumerate(cos):
        txt = co.get("text", "") if isinstance(co, dict) else str(co)
        note = co.get("note") if isinstance(co, dict) else None
        tw_ = cw - bub - Pt(12)
        _numbubble(sl, cx, y, co.get("n", i + 1) if isinstance(co, dict) else i + 1,
                   d=bub)
        _text(sl, cx + bub + Pt(12), y + Pt(1), tw_, Pt(46),
              txt, size=size, color=T.GRAY, spacing=1.2)
        y += gb + gl * (_wrap_lines(txt, tw_, fs) - 1)
        if note:
            _text(sl, cx + bub + Pt(12), y - Pt(4), tw_, Pt(38),
                  "⚠ " + note, size=Pt(max(fs - 2, 9)), color=T.RED,
                  bold=True, spacing=1.2)
            y += Pt(28) + Pt(17) * (_wrap_lines(note, tw_, fs - 2) - 1)
    return sl


def lay_screenshot_bullets(prs, s):
    """L6 - Ảnh + gạch đầu dòng (không đánh số)."""
    sl = _blank(prs)
    _title(sl, s["title"], s.get("subtitle"))
    side = s.get("image_side", "left")
    imgw, imgh = Pt(510), Pt(348)
    tw = T.SLIDE_W - T.M_LEFT * 2 - imgw - Pt(28)
    if side == "right":
        tx, ix = T.M_LEFT, T.M_LEFT + tw + Pt(28)
    else:
        ix, tx = T.M_LEFT, T.M_LEFT + imgw + Pt(28)
    _image(sl, s.get("image"), ix, T.CONTENT_TOP, imgw, imgh)
    _bullets(sl, tx, T.CONTENT_TOP + Pt(6), tw, imgh,
             s.get("bullets", []), size=T.SZ_CALLOUT)
    if s.get("note"):
        _rect(sl, tx, T.CONTENT_TOP + Pt(268), tw, Pt(80), fill=T.LIGHT)
        _text(sl, tx + Pt(14), T.CONTENT_TOP + Pt(279), tw - Pt(28), Pt(62),
              "⚠ " + s["note"], size=T.sz(12), color=T.RED, bold=True, spacing=1.2)
    return sl


def lay_process(prs, s):
    """L7 - Sơ đồ luồng ngang các giai đoạn, có SLA và ghi chú dưới mỗi bước."""
    sl = _blank(prs)
    _title(sl, s["title"], s.get("subtitle"))
    steps = s.get("steps", [])
    n = max(len(steps), 1)
    gap = Pt(12)
    total = T.SLIDE_W - T.M_LEFT * 2
    w = (total - gap * (n - 1)) / n
    ytop = T.CONTENT_TOP + Pt(40)

    for i, st in enumerate(steps):
        x = T.M_LEFT + i * (w + gap)
        hi = st.get("highlight")          # bước do Store làm -> tô đỏ
        fill = T.RED if hi else T.LIGHT
        fg = T.WHITE if hi else T.GRAY

        _rect(sl, x, ytop, w, Pt(86), fill=fill, line=None if hi else T.BORDER)
        _text(sl, x + Pt(6), ytop + Pt(10), w - Pt(12), Pt(66),
              st.get("name", ""), size=T.sz(14), color=fg, bold=True,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.1)

        if st.get("sla"):
            _text(sl, x, ytop - Pt(30), w, Pt(24), st["sla"],
                  size=T.sz(12), color=T.RED, bold=True, align=PP_ALIGN.CENTER)
        if st.get("actor"):
            _text(sl, x, ytop + Pt(94), w, Pt(24), st["actor"],
                  size=T.sz(11), color=T.YELLOW if hi else T.BORDER,
                  bold=True, align=PP_ALIGN.CENTER)
        if st.get("desc"):
            _text(sl, x + Pt(2), ytop + Pt(122), w - Pt(4), Pt(150),
                  st["desc"], size=T.sz(11), color=T.GRAY,
                  align=PP_ALIGN.CENTER, spacing=1.2)
        if i < n - 1:
            _text(sl, x + w - Pt(2), ytop + Pt(28), gap + Pt(4), Pt(32),
                  "›", size=T.sz(20), color=T.RED, bold=True, align=PP_ALIGN.CENTER)

    if s.get("legend"):
        _text(sl, T.M_LEFT, T.SLIDE_H - Pt(56), total, Pt(24),
              s["legend"], size=T.sz(12), color=T.GRAY)
    return sl


def lay_table(prs, s):
    """Bảng tóm tắt - vd 'Store chỉ cần nhớ 3 nút'."""
    sl = _blank(prs)
    _title(sl, s["title"], s.get("subtitle"))
    hdr = s.get("headers", [])
    rows = s.get("rows", [])
    nr, nc = len(rows) + 1, max(len(hdr), 1)
    h = min(Pt(62) * nr, Pt(348))
    shape = sl.shapes.add_table(nr, nc, T.M_LEFT, T.CONTENT_TOP,
                                T.SLIDE_W - T.M_LEFT * 2, h)
    tbl = shape.table
    if s.get("col_widths"):
        tot = sum(s["col_widths"])
        avail = T.SLIDE_W - T.M_LEFT * 2
        for i, cwd in enumerate(s["col_widths"]):
            tbl.columns[i].width = Emu(int(avail * cwd / tot))

    for c, txt in enumerate(hdr):
        cell = tbl.cell(0, c)
        cell.text = str(txt)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _c(T.RED)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = T.sz(14); r.font.bold = True
                r.font.color.rgb = _c(T.WHITE); r.font.name = T.FONT

    for ri, row in enumerate(rows, start=1):
        for c in range(nc):
            cell = tbl.cell(ri, c)
            cell.text = str(row[c]) if c < len(row) else ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = _c(T.WHITE if ri % 2 else T.LIGHT)
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = T.sz(13)
                    r.font.color.rgb = _c(T.GRAY); r.font.name = T.FONT
    return sl


LAYOUTS = {
    "title": lay_title,
    "section": lay_section,
    "agenda": lay_agenda,
    "two_col": lay_two_col,
    "screenshot_callout": lay_screenshot_callout,
    "screenshot_bullets": lay_screenshot_bullets,
    "process": lay_process,
    "table": lay_table,
}


# ============================== chạy ==============================

def build(data, outpath):
    prs = Presentation()
    prs.slide_width, prs.slide_height = T.SLIDE_W, T.SLIDE_H

    slides = data.get("slides", [])
    missing = []
    for i, s in enumerate(slides, start=1):
        fn = LAYOUTS.get(s.get("layout"))
        if not fn:
            print(f"  ! slide {i}: layout '{s.get('layout')}' không tồn tại - bỏ qua")
            continue
        sl = fn(prs, s)
        if i > 1:
            _footer(sl, i, len(slides))
        # logo góc chỉ đóng lên slide nội dung, không đóng lên bìa/chương
        if s.get("layout") not in ("title", "section"):
            _logo_mark(sl)
        img = s.get("image")
        if img and not os.path.exists(img):
            missing.append((i, img))

    os.makedirs(os.path.dirname(outpath) or ".", exist_ok=True)
    prs.save(outpath)
    return len(slides), missing


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else "content.yaml"
    out = sys.argv[2] if len(sys.argv) > 2 else "out/deck.pptx"

    with open(src, encoding="utf-8") as f:
        data = yaml_or_json(f.read(), src)

    n, missing = build(data, out)
    print(f"[OK] {n} slide -> {out}")
    if missing:
        print(f"\n[!] {len(missing)} ảnh chưa có (slide hiện ô gạch chờ):")
        for i, p in missing:
            print(f"    slide {i:>2}: {p}")
        print(f"\n    Xem cần chụp những gì:  python check_shots.py --todo")


def yaml_or_json(text, name):
    if name.lower().endswith((".yaml", ".yml")):
        import yaml
        return yaml.safe_load(text)
    return json.loads(text)


if __name__ == "__main__":
    main()
